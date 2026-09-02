import pandas as pd
from openpyxl.styles import PatternFill

# ----------------配置区域----------------
file_path = "data.xlsx"          # 你的输入 Excel 文件路径
sheet_name = 0                   # 工作表名称或索引（0 表示第一个 sheet）
output_file = "diff_result.xlsx" # 导出结果文件名
# ----------------------------------------

def compare_excel_columns():
    # 1. 读取 Excel 文件（不将第一行作为表头，确保能按位置读取）
    df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
    
    # 获取前两列数据并转为字符串处理（防止格式不一致导致比较出错）
    col_a_data = df.iloc[:, 0].astype(str).str.strip()
    col_b_data = df.iloc[:, 1].astype(str).str.strip()
    
    # 2. 集合对比（获取 A 有 B 无的数据列表）
    set_a = set(col_a_data.dropna())
    set_b = set(col_b_data.dropna())
    
    # 过滤掉空的字符串标识
    set_a.discard("nan")
    set_b.discard("nan")
    
    only_in_a = list(set_a - set_b)
    only_in_b = list(set_b - set_a)

    print("=== 【对比统计】 ===")
    print(f"B列缺失的数据数量（A有B无）: {len(only_in_a)}")
    print(f"A列缺失的数据数量（B有A无）: {len(only_in_b)}")

    # 3. 逐行对比生成明细
    max_len = max(len(col_a_data), len(col_b_data))
    row_diff = []
    
    for idx in range(max_len):
        val_a = col_a_data.iloc[idx] if idx < len(col_a_data) else "nan"
        val_b = col_b_data.iloc[idx] if idx < len(col_b_data) else "nan"
        
        if val_a != val_b:
            row_diff.append({
                "Excel行号": idx + 1,
                "A列数据": val_a if val_a != "nan" else "[空]",
                "B列数据": val_b if val_b != "nan" else "[空]"
            })

    # 4. 导出为新的 Excel 文件并为缺失项高亮标记红色
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        # 保存对比统计结果页签
        pd.DataFrame({"B列缺失数据(A有B无)": only_in_a}).to_excel(writer, sheet_name="B列缺失数据", index=False)
        pd.DataFrame({"A列缺失数据(B有A无)": only_in_b}).to_excel(writer, sheet_name="A列缺失数据", index=False)
        pd.DataFrame(row_diff).to_excel(writer, sheet_name="逐行差异明细", index=False)
        
        # 将原表格完整导出，并在 A 列中将 B 列缺少的数据标红
        df_export = df.copy()
        # 加上默认列名
        df_export.columns = [f"列_{i+1}" for i in range(df_export.shape[1])]
        df_export.to_excel(writer, sheet_name="原表标记结果", index=False)
        
        # 获取 openpyxl 的 worksheet 对象进行颜色标记
        ws = writer.sheets["原表标记结果"]
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # 淡红色填充
        
        # 遍历 A 列单元格，如果数据存在于 only_in_a 中则标记为红色
        for row_idx, val in enumerate(col_a_data, start=2): # 因为导出的 excel 带了一行生成的表头，所以从第 2 行开始
            if val in only_in_a:
                ws.cell(row=row_idx, column=1).fill = red_fill

    print(f"\n对比完成！结果文件已成功生成: {output_file}")

if __name__ == "__main__":
    compare_excel_columns()